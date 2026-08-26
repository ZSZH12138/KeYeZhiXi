"""Private deterministic byte parsers used by M1 course governance.

This module deliberately accepts captured bytes rather than filesystem paths so
the bytes hashed, parsed, and later persisted by M1 are identical.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePath
from typing import Protocol
import unicodedata
from zipfile import BadZipFile, ZipFile

from course_insight.modules.m1_course_governance.parser_protocol import (
    DEFAULT_MAX_BYTES,
    ParserEntry,
    ParserRegistry,
)
from course_insight.modules.m1_course_governance.legacy_powerpoint import (
    LegacyPowerPointConversionError,
    LegacyPowerPointConverter,
)

_MAX_RAW_BYTES = DEFAULT_MAX_BYTES
_MAX_PDF_PAGES = 200
_MAX_PPTX_SLIDES = 200
_MAX_OOXML_MEMBERS = 4_096
_MAX_OOXML_UNCOMPRESSED_BYTES = 16 * 1024 * 1024
_MAX_PDF_PAGE_TEXT_BYTES = 2 * 1024 * 1024
_MAX_PDF_TOTAL_TEXT_BYTES = 8 * 1024 * 1024


class _ParserError(ValueError):
    """A safe, private parser failure reason for the M1 service to map."""


class OCRTextProvider(Protocol):
    """Injectable OCR adapter for image-only PDF pages.

    The parser owns byte limits and page locators; an adapter owns the
    environment-specific PDF rasterisation and OCR executable.
    """

    def extract_text(self, payload: bytes, *, page_number: int) -> str:
        """Return text for one 1-based PDF page."""


class OCRRequiredError(_ParserError):
    """The PDF has no text layer and no OCR provider was configured."""

    code = "COURSE_PDF_OCR_REQUIRED"


@dataclass(frozen=True, slots=True)
class ParsedBlock:
    """One normalized text block with a stable source-relative locator."""

    text: str
    locator: str
    ordinal: int


@dataclass(frozen=True, slots=True)
class ParsedSource:
    """Private parsing result; it is not a public Pydantic contract."""

    media_type: str
    page_count: int | None
    blocks: tuple[ParsedBlock, ...]


def parse_source(
    file_name: str,
    payload: bytes,
    *,
    ocr_provider: OCRTextProvider | None = None,
    legacy_powerpoint_converter: LegacyPowerPointConverter | None = None,
) -> ParsedSource:
    """Parse one supported source byte sequence without accessing the host."""

    if not isinstance(payload, bytes):
        raise _ParserError("source payload must be bytes")
    if len(payload) > _MAX_RAW_BYTES:
        raise _ParserError("source exceeds raw byte limit")

    suffix = PurePath(file_name).suffix.casefold()
    if suffix == ".ppt":
        if legacy_powerpoint_converter is None:
            raise _ParserError("legacy PowerPoint conversion is unavailable")
        try:
            converted = legacy_powerpoint_converter.convert(payload)
        except LegacyPowerPointConversionError:
            raise
        except Exception as error:
            raise _ParserError("legacy PowerPoint conversion failed") from error
        parsed = _parse_pptx(converted)
        return ParsedSource(
            media_type="application/vnd.ms-powerpoint",
            page_count=parsed.page_count,
            blocks=parsed.blocks,
        )
    if suffix == ".md":
        return _parse_plain_text(payload, media_type="text/markdown")
    if suffix == ".txt":
        return _parse_plain_text(payload, media_type="text/plain")
    if suffix == ".pdf":
        return _parse_pdf(payload, ocr_provider=ocr_provider)
    if suffix == ".docx":
        return _parse_docx(payload)
    if suffix == ".pptx":
        return _parse_pptx(payload)
    raise _ParserError("unsupported source file type")


def default_parser_registry(
    *,
    ocr_provider: OCRTextProvider | None = None,
    legacy_powerpoint_converter: LegacyPowerPointConverter | None = None,
) -> ParserRegistry:
    """Build the versioned registry for the built-in M1 byte parsers."""

    pdf_parser = (
        parse_source
        if ocr_provider is None
        else lambda file_name, payload: parse_source(
            file_name,
            payload,
            ocr_provider=ocr_provider,
        )
    )
    pdf_capabilities = frozenset({"byte-input", "paged"})
    if ocr_provider is not None:
        pdf_capabilities = frozenset({"byte-input", "paged", "ocr"})
    entries = [
            ParserEntry(
                extension=".md",
                media_type="text/markdown",
                parser_version="parse-source-v1",
                capabilities=frozenset({"byte-input", "text"}),
                max_bytes=_MAX_RAW_BYTES,
                parser=parse_source,
                parser_id="m1.parse_source",
            ),
            ParserEntry(
                extension=".txt",
                media_type="text/plain",
                parser_version="parse-source-v1",
                capabilities=frozenset({"byte-input", "text"}),
                max_bytes=_MAX_RAW_BYTES,
                parser=parse_source,
                parser_id="m1.parse_source",
            ),
            ParserEntry(
                extension=".pdf",
                media_type="application/pdf",
                parser_version="parse-source-v1",
                capabilities=pdf_capabilities,
                max_bytes=_MAX_RAW_BYTES,
                parser=pdf_parser,
                parser_id=(
                    "m1.parse_source"
                    if ocr_provider is None
                    else "m1.parse_source.ocr"
                ),
            ),
            ParserEntry(
                extension=".docx",
                media_type=(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                ),
                parser_version="parse-source-v1",
                capabilities=frozenset({"byte-input", "structured"}),
                max_bytes=_MAX_RAW_BYTES,
                parser=parse_source,
                parser_id="m1.parse_source",
            ),
            ParserEntry(
                extension=".pptx",
                media_type=(
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                ),
                parser_version="parse-source-v1",
                capabilities=frozenset({"byte-input", "paged", "structured"}),
                max_bytes=_MAX_RAW_BYTES,
                parser=parse_source,
                parser_id="m1.parse_source",
            ),
    ]
    if legacy_powerpoint_converter is not None:
        entries.append(
            ParserEntry(
                extension=".ppt",
                media_type="application/vnd.ms-powerpoint",
                parser_version="parse-source-v1",
                capabilities=frozenset(
                    {"byte-input", "converted", "paged", "structured"}
                ),
                max_bytes=_MAX_RAW_BYTES,
                parser=lambda file_name, payload: parse_source(
                    file_name,
                    payload,
                    legacy_powerpoint_converter=legacy_powerpoint_converter,
                ),
                parser_id="m1.parse_source.powerpoint-com",
            )
        )
    return ParserRegistry(entries)


def _parse_plain_text(payload: bytes, *, media_type: str) -> ParsedSource:
    try:
        decoded = payload.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise _ParserError("text source is not valid utf-8") from error
    lines = [
        unicodedata.normalize("NFKC", line).rstrip()
        for line in decoded.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    ]
    blocks: list[ParsedBlock] = []
    block_lines: list[str] = []
    first_line: int | None = None

    def append_block(last_line: int) -> None:
        nonlocal block_lines, first_line
        if first_line is None:
            return
        text = "\n".join(block_lines).strip()
        if text:
            ordinal = len(blocks) + 1
            blocks.append(
                ParsedBlock(
                    text=text,
                    locator=(
                        f"paragraph:{ordinal};lines:{first_line}-{last_line}"
                    ),
                    ordinal=ordinal,
                )
            )
        block_lines = []
        first_line = None

    for line_number, line in enumerate(lines, start=1):
        if line.strip():
            if first_line is None:
                first_line = line_number
            block_lines.append(line)
        else:
            append_block(line_number - 1)
    append_block(len(lines))
    if not blocks:
        raise _ParserError("source contains no text blocks")
    return ParsedSource(media_type=media_type, page_count=None, blocks=tuple(blocks))


def _parse_pdf(
    payload: bytes,
    *,
    ocr_provider: OCRTextProvider | None = None,
) -> ParsedSource:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(payload), strict=True)
        if reader.is_encrypted:
            raise _ParserError("encrypted PDF sources are not supported")
        page_count = len(reader.pages)
        if page_count > _MAX_PDF_PAGES:
            raise _ParserError("PDF exceeds page limit")
        blocks: list[ParsedBlock] = []
        total_text_bytes = 0
        for page_number, page in enumerate(reader.pages, start=1):
            extracted_text = _normalize_text(page.extract_text() or "")
            if not extracted_text and ocr_provider is not None:
                extracted_text = _normalize_text(
                    ocr_provider.extract_text(
                        payload,
                        page_number=page_number,
                    )
                )
            page_text_bytes = len(extracted_text.encode("utf-8"))
            if page_text_bytes > _MAX_PDF_PAGE_TEXT_BYTES:
                raise _ParserError("PDF extracted text byte limit exceeded")
            total_text_bytes += page_text_bytes
            if total_text_bytes > _MAX_PDF_TOTAL_TEXT_BYTES:
                raise _ParserError("PDF extracted text byte limit exceeded")
            _append_paragraph_blocks(
                blocks,
                extracted_text,
                locator_prefix=f"page:{page_number};block",
            )
    except _ParserError:
        raise
    except Exception as error:
        raise _ParserError("PDF source could not be parsed") from error
    if not blocks:
        raise OCRRequiredError(
            "PDF has no text layer; configure an OCR provider before import"
        )
    return ParsedSource(
        media_type="application/pdf",
        page_count=page_count,
        blocks=tuple(blocks),
    )


def _parse_docx(payload: bytes) -> ParsedSource:
    _validate_ooxml_zip(payload)
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import _Cell
    from docx.text.paragraph import Paragraph

    try:
        document = Document(BytesIO(payload))
        blocks: list[ParsedBlock] = []
        paragraph_number = 0
        table_number = 0
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                paragraph_number += 1
                _append_block(
                    blocks,
                    Paragraph(child, document).text,
                    locator=f"paragraph:{paragraph_number}",
                )
            elif child.tag == qn("w:tbl"):
                table_number += 1
                _append_docx_table(
                    blocks,
                    child,
                    parent=document,
                    table_number=table_number,
                    cell_type=_Cell,
                )
    except _ParserError:
        raise
    except Exception as error:
        raise _ParserError("DOCX source could not be parsed") from error
    if not blocks:
        raise _ParserError("source contains no text blocks")
    return ParsedSource(
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        page_count=None,
        blocks=tuple(blocks),
    )


def _append_docx_table(
    blocks: list[ParsedBlock],
    table_xml: object,
    *,
    parent: object,
    table_number: int,
    cell_type: object,
) -> None:
    for row_number, row_xml in enumerate(table_xml.tr_lst, start=1):
        cell_number = 1
        for cell_xml in row_xml.tc_lst:
            properties = cell_xml.tcPr
            grid_span = (
                1
                if properties.gridSpan is None
                else int(properties.gridSpan.val)
            )
            vertical_merge = properties.vMerge
            if (
                vertical_merge is not None
                and vertical_merge.val != "restart"
            ):
                cell_number += grid_span
                continue
            cell = cell_type(cell_xml, parent)
            _append_block(
                blocks,
                cell.text,
                locator=(
                    f"table:{table_number};row:{row_number};cell:{cell_number}"
                ),
            )
            cell_number += grid_span


def _parse_pptx(payload: bytes) -> ParsedSource:
    _validate_ooxml_zip(payload)
    from pptx import Presentation

    try:
        presentation = Presentation(BytesIO(payload))
        page_count = len(presentation.slides)
        if page_count > _MAX_PPTX_SLIDES:
            raise _ParserError("PPTX exceeds slide limit")
        blocks: list[ParsedBlock] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            for shape_number, shape in enumerate(
                _iter_shape_tree(slide.shapes), start=1
            ):
                if getattr(shape, "has_text_frame", False):
                    for paragraph_number, paragraph in enumerate(
                        shape.text_frame.paragraphs,
                        start=1,
                    ):
                        _append_block(
                            blocks,
                            paragraph.text,
                            locator=(
                                f"slide:{slide_number};shape:{shape_number};"
                                f"paragraph:{paragraph_number}"
                            ),
                        )
                if getattr(shape, "has_table", False):
                    _append_pptx_table(
                        blocks,
                        shape.table,
                        slide_number=slide_number,
                        shape_number=shape_number,
                    )
    except _ParserError:
        raise
    except Exception as error:
        raise _ParserError("PPTX source could not be parsed") from error
    if not blocks:
        raise _ParserError("source contains no text blocks")
    return ParsedSource(
        media_type=(
            "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        ),
        page_count=page_count,
        blocks=tuple(blocks),
    )


def _iter_shape_tree(shapes: object):
    for shape in shapes:
        yield shape
        child_shapes = getattr(shape, "shapes", None)
        if child_shapes is not None:
            yield from _iter_shape_tree(child_shapes)


def _append_pptx_table(
    blocks: list[ParsedBlock],
    table: object,
    *,
    slide_number: int,
    shape_number: int,
) -> None:
    for row_number, row in enumerate(table.rows, start=1):
        for cell_number, cell in enumerate(row.cells, start=1):
            if cell.is_spanned:
                continue
            _append_block(
                blocks,
                cell.text,
                locator=(
                    f"slide:{slide_number};table:{shape_number};"
                    f"row:{row_number};cell:{cell_number}"
                ),
            )


def _validate_ooxml_zip(payload: bytes) -> None:
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members = archive.infolist()
    except (BadZipFile, OSError) as error:
        raise _ParserError("OOXML ZIP source could not be parsed") from error
    if len(members) > _MAX_OOXML_MEMBERS:
        raise _ParserError("OOXML ZIP exceeds member limit")
    total_uncompressed = 0
    for member in members:
        if member.flag_bits & 0x1:
            raise _ParserError("encrypted OOXML ZIP sources are not supported")
        total_uncompressed += member.file_size
        if total_uncompressed > _MAX_OOXML_UNCOMPRESSED_BYTES:
            raise _ParserError("OOXML ZIP exceeds uncompressed byte limit")


def _append_paragraph_blocks(
    blocks: list[ParsedBlock],
    text: str,
    *,
    locator_prefix: str,
) -> None:
    lines = _normalized_lines(text)
    paragraph_lines: list[str] = []
    for line in lines:
        if line.strip():
            paragraph_lines.append(line)
        elif paragraph_lines:
            _append_block(
                blocks,
                "\n".join(paragraph_lines),
                locator=f"{locator_prefix}:{_next_block_number(blocks, locator_prefix)}",
            )
            paragraph_lines = []
    if paragraph_lines:
        _append_block(
            blocks,
            "\n".join(paragraph_lines),
            locator=f"{locator_prefix}:{_next_block_number(blocks, locator_prefix)}",
        )


def _next_block_number(blocks: list[ParsedBlock], locator_prefix: str) -> int:
    return 1 + sum(
        1 for block in blocks if block.locator.startswith(f"{locator_prefix}:")
    )


def _append_block(
    blocks: list[ParsedBlock],
    text: str,
    *,
    locator: str,
) -> None:
    normalized = _normalize_text(text)
    if normalized:
        blocks.append(
            ParsedBlock(
                text=normalized,
                locator=locator,
                ordinal=len(blocks) + 1,
            )
        )


def _normalized_lines(text: str) -> list[str]:
    return [
        unicodedata.normalize("NFKC", line).rstrip()
        for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    ]


def _normalize_text(text: str) -> str:
    return "\n".join(_normalized_lines(text)).strip()
