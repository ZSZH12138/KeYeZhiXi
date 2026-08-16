"""Deterministic, byte-oriented M1 parser boundary tests."""

from __future__ import annotations

import importlib
from io import BytesIO
import zlib
from zipfile import ZIP_DEFLATED, ZipFile

import pytest

from course_insight.modules.m1_course_governance.parsers import (
    _ParserError,
    parse_source,
)


def test_markdown_normalizes_text_and_retains_source_line_ranges() -> None:
    parsed = parse_source(
        "lesson.MD",
        "\ufeffAlpha\r\nBeta  \r\n\r\n\uff23\uff4f\uff55\uff52\uff53\uff45\r\n".encode("utf-8"),
    )

    assert parsed.media_type == "text/markdown"
    assert parsed.page_count is None
    assert [(block.text, block.locator, block.ordinal) for block in parsed.blocks] == [
        ("Alpha\nBeta", "paragraph:1;lines:1-2", 1),
        ("Course", "paragraph:2;lines:4-4", 2),
    ]


def test_plain_text_drops_empty_paragraphs_and_uses_plain_text_media_type() -> None:
    parsed = parse_source("notes.txt", b"\n  \nOne\n\n\nTwo\n")

    assert parsed.media_type == "text/plain"
    assert [(block.text, block.locator, block.ordinal) for block in parsed.blocks] == [
        ("One", "paragraph:1;lines:3-3", 1),
        ("Two", "paragraph:2;lines:6-6", 2),
    ]


def test_legacy_powerpoint_is_explicitly_rejected() -> None:
    with pytest.raises(_ParserError, match="unsupported legacy powerpoint"):
        parse_source("lecture.ppt", b"not a modern presentation")


def test_empty_text_source_is_rejected() -> None:
    with pytest.raises(_ParserError, match="source contains no text blocks"):
        parse_source("notes.txt", b" \r\n\r\n")


def test_invalid_utf8_text_source_is_rejected() -> None:
    with pytest.raises(_ParserError, match="not valid utf-8"):
        parse_source("notes.txt", b"\xff\xfe")


def test_raw_source_limit_fails_closed() -> None:
    payload = b"x" * (25 * 1024 * 1024 + 1)

    with pytest.raises(_ParserError, match="raw byte limit"):
        parse_source("notes.txt", payload)


def test_pdf_extracts_one_normalized_block_per_nonempty_text_block() -> None:
    parsed = parse_source("lesson.pdf", _single_page_pdf("AlphaBeta"))

    assert parsed.media_type == "application/pdf"
    assert parsed.page_count == 1
    assert [(block.text, block.locator, block.ordinal) for block in parsed.blocks] == [
        ("AlphaBeta", "page:1;block:1", 1),
    ]


def test_binary_parser_dependencies_are_available() -> None:
    assert importlib.import_module("pypdf") is not None
    assert importlib.import_module("docx") is not None
    assert importlib.import_module("pptx") is not None


def test_pdf_rejects_encrypted_and_corrupt_sources_without_host_path_details() -> None:
    from pypdf import PdfWriter

    encrypted = PdfWriter()
    encrypted.add_blank_page(width=72, height=72)
    encrypted.encrypt("secret")
    encrypted_payload = BytesIO()
    encrypted.write(encrypted_payload)

    for payload in (encrypted_payload.getvalue(), b"not a PDF"):
        with pytest.raises(_ParserError) as error:
            parse_source("lecture.pdf", payload)
        assert "lecture.pdf" not in str(error.value)


def test_pdf_page_limit_fails_closed() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    for _ in range(201):
        writer.add_blank_page(width=72, height=72)
    payload = BytesIO()
    writer.write(payload)

    with pytest.raises(_ParserError, match="page limit"):
        parse_source("many-pages.pdf", payload.getvalue())


def test_pdf_extracted_text_limit_fails_closed_for_small_flate_payload() -> None:
    payload = _single_page_pdf("A" * 4_000_000, flate=True)
    assert len(payload) < 5_000

    with pytest.raises(_ParserError, match="extracted text byte limit"):
        parse_source("expanded.pdf", payload)


def test_pdf_cumulative_extracted_text_limit_fails_closed() -> None:
    payload = _multi_page_flate_pdf(text="A" * 1_700_000, page_count=5)
    assert len(payload) < 20_000

    with pytest.raises(_ParserError, match="extracted text byte limit"):
        parse_source("cumulative-expanded.pdf", payload)


def test_docx_paragraphs_and_table_cells_keep_document_order() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("First")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Second"
    table.cell(0, 1).text = "\uff34\uff48\uff49\uff52\uff44"
    document.add_paragraph("Fourth")
    payload = BytesIO()
    document.save(payload)

    parsed = parse_source("lesson.docx", payload.getvalue())

    assert parsed.media_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert parsed.page_count is None
    assert [(block.text, block.locator) for block in parsed.blocks] == [
        ("First", "paragraph:1"),
        ("Second", "table:1;row:1;cell:1"),
        ("Third", "table:1;row:1;cell:2"),
        ("Fourth", "paragraph:2"),
    ]


def test_docx_merged_cells_emit_only_visible_roots_with_physical_grid_columns() -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "Vertical"
    table.cell(0, 1).merge(table.cell(1, 2)).text = "Rectangle"
    table.cell(2, 0).text = "Bottom one"
    table.cell(2, 1).text = "Bottom two"
    table.cell(2, 2).text = "Bottom three"
    payload = BytesIO()
    document.save(payload)

    parsed = parse_source("merged.docx", payload.getvalue())

    assert [(block.text, block.locator) for block in parsed.blocks] == [
        ("Vertical", "table:1;row:1;cell:1"),
        ("Rectangle", "table:1;row:1;cell:2"),
        ("Bottom one", "table:1;row:3;cell:1"),
        ("Bottom two", "table:1;row:3;cell:2"),
        ("Bottom three", "table:1;row:3;cell:3"),
    ]


def test_valid_highly_compressible_docx_is_not_rejected_by_ratio_alone() -> None:
    from docx import Document

    document = Document()
    document.add_paragraph("A" * 100_000)
    payload = BytesIO()
    document.save(payload)

    parsed = parse_source("repetitive.docx", payload.getvalue())

    assert parsed.blocks[0].text == "A" * 100_000


def test_pptx_paragraphs_and_visible_table_cells_follow_shape_tree_order() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    text_box.text_frame.text = "First"
    text_box.text_frame.add_paragraph().text = "\uff33econd"
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(4), Inches(1))
    table_shape.table.cell(0, 0).text = "Third"
    table_shape.table.cell(0, 1).text = "Fourth"
    payload = BytesIO()
    presentation.save(payload)

    parsed = parse_source("lesson.pptx", payload.getvalue())

    assert parsed.media_type == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert parsed.page_count == 1
    assert [(block.text, block.locator) for block in parsed.blocks] == [
        ("First", "slide:1;shape:1;paragraph:1"),
        ("Second", "slide:1;shape:1;paragraph:2"),
        ("Third", "slide:1;table:2;row:1;cell:1"),
        ("Fourth", "slide:1;table:2;row:1;cell:2"),
    ]


def test_pptx_merged_table_cells_are_emitted_once_at_the_visible_origin() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged"
    payload = BytesIO()
    presentation.save(payload)

    parsed = parse_source("merged.pptx", payload.getvalue())

    assert [(block.text, block.locator) for block in parsed.blocks] == [
        ("Merged", "slide:1;table:1;row:1;cell:1"),
    ]


def test_pptx_slide_limit_fails_closed() -> None:
    from pptx import Presentation

    presentation = Presentation()
    for _ in range(201):
        presentation.slides.add_slide(presentation.slide_layouts[6])
    payload = BytesIO()
    presentation.save(payload)

    with pytest.raises(_ParserError, match="slide limit"):
        parse_source("many-slides.pptx", payload.getvalue())


def test_ooxml_zip_bomb_fails_before_document_parser_consumes_it() -> None:
    payload = BytesIO()
    with ZipFile(payload, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * (16 * 1024 * 1024 + 1))

    with pytest.raises(_ParserError, match="OOXML ZIP"):
        parse_source("bomb.docx", payload.getvalue())


def test_corrupt_ooxml_source_is_rejected() -> None:
    with pytest.raises(_ParserError, match="OOXML ZIP"):
        parse_source("broken.pptx", b"not a ZIP file")


def _single_page_pdf(text: str, *, flate: bool = False) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("utf-8")
    encoded_stream = zlib.compress(stream) if flate else stream
    filter_part = b" /Filter /FlateDecode" if flate else b""
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
            b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        (
            b"<< /Length "
            + str(len(encoded_stream)).encode("ascii")
            + filter_part
            + b" >>\nstream\n"
            + encoded_stream
            + b"\nendstream"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    )
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, value in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{number} 0 obj\n".encode("ascii"))
        payload.extend(value)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(payload)


def _multi_page_flate_pdf(*, text: str, page_count: int) -> bytes:
    from pypdf import PdfReader, PdfWriter

    source = PdfReader(BytesIO(_single_page_pdf(text, flate=True)))
    writer = PdfWriter()
    for _ in range(page_count):
        writer.add_page(source.pages[0])
    payload = BytesIO()
    writer.write(payload)
    return payload.getvalue()
