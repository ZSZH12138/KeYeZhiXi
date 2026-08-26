from __future__ import annotations

import base64
import subprocess
from io import BytesIO
from pathlib import Path

import pytest

from course_insight.modules.m1_course_governance.legacy_powerpoint import (
    LegacyPowerPointConversionFailed,
    LegacyPowerPointConversionUnavailable,
    PowerPointComConverter,
)


_OLE_COMPOUND_FILE_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


def _pptx_bytes(text: str = "拥塞控制") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = text
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def test_converter_uses_hardened_automation_and_cleans_temporary_files(
    monkeypatch,
) -> None:
    converted = _pptx_bytes()
    observed: dict[str, object] = {}

    def runner(command, **options):
        environment = options["env"]
        input_path = Path(environment["COURSE_INSIGHT_PPT_INPUT"])
        output_path = Path(environment["COURSE_INSIGHT_PPT_OUTPUT"])
        observed.update(
            command=command,
            options=options,
            input_path=input_path,
            output_path=output_path,
            input_bytes=input_path.read_bytes(),
        )
        output_path.write_bytes(converted)
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    monkeypatch.setenv("DEEPSEEK_API_KEY", "must-not-reach-converter")
    source = _OLE_COMPOUND_FILE_SIGNATURE + b"legacy presentation"
    result = PowerPointComConverter(
        powershell_executable="powershell.exe",
        runner=runner,
    ).convert(source)

    command = observed["command"]
    options = observed["options"]
    encoded_script = command[command.index("-EncodedCommand") + 1]
    script = base64.b64decode(encoded_script).decode("utf-16-le")
    assert result == converted
    assert observed["input_bytes"] == source
    assert options["timeout"] == 120.0
    assert options["capture_output"] is True
    assert options["text"] is True
    assert "DEEPSEEK_API_KEY" not in options["env"]
    assert "AutomationSecurity = 3" in script
    assert "Presentations.Open($inputPath, -1, 0, 0)" in script
    assert "SaveAs($outputPath, 24)" in script
    assert not observed["input_path"].exists()
    assert not observed["output_path"].exists()


def test_converter_rejects_non_ole_bytes_before_launching_powerpoint() -> None:
    def runner(command, **options):
        del command, options
        raise AssertionError("invalid input must not launch PowerPoint")

    with pytest.raises(LegacyPowerPointConversionFailed) as raised:
        PowerPointComConverter(
            powershell_executable="powershell.exe",
            runner=runner,
        ).convert(b"renamed text file")

    assert raised.value.code == "LEGACY_PPT_CONVERSION_FAILED"


def test_converter_reports_powerpoint_unavailable_without_process_output() -> None:
    def runner(command, **options):
        return subprocess.CompletedProcess(
            command,
            42,
            stdout="private stdout",
            stderr="private host details",
        )

    with pytest.raises(LegacyPowerPointConversionUnavailable) as raised:
        PowerPointComConverter(
            powershell_executable="powershell.exe",
            runner=runner,
        ).convert(_OLE_COMPOUND_FILE_SIGNATURE + b"legacy")

    assert raised.value.code == "LEGACY_PPT_CONVERSION_UNAVAILABLE"
    assert "private" not in str(raised.value)


def test_converter_maps_timeout_to_safe_failure() -> None:
    def runner(command, **options):
        raise subprocess.TimeoutExpired(command, options["timeout"])

    with pytest.raises(LegacyPowerPointConversionFailed, match="timed out"):
        PowerPointComConverter(
            powershell_executable="powershell.exe",
            runner=runner,
        ).convert(_OLE_COMPOUND_FILE_SIGNATURE + b"legacy")


@pytest.mark.parametrize("output", [None, b"", b"x" * 33])
def test_converter_requires_nonempty_bounded_output(output: bytes | None) -> None:
    def runner(command, **options):
        if output is not None:
            Path(options["env"]["COURSE_INSIGHT_PPT_OUTPUT"]).write_bytes(output)
        return subprocess.CompletedProcess(command, 0, stdout="", stderr="")

    with pytest.raises(LegacyPowerPointConversionFailed, match="valid output"):
        PowerPointComConverter(
            powershell_executable="powershell.exe",
            runner=runner,
            max_output_bytes=32,
        ).convert(_OLE_COMPOUND_FILE_SIGNATURE + b"legacy")


def test_converter_reports_missing_powershell_as_unavailable() -> None:
    def runner(command, **options):
        del command, options
        raise FileNotFoundError

    with pytest.raises(LegacyPowerPointConversionUnavailable):
        PowerPointComConverter(
            powershell_executable="missing-powershell.exe",
            runner=runner,
        ).convert(_OLE_COMPOUND_FILE_SIGNATURE + b"legacy")
