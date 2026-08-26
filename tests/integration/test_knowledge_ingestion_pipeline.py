from __future__ import annotations

import hashlib
import logging
import threading
import time
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest

from course_insight.application.knowledge_ingestion import KnowledgeIngestionProcessor
from course_insight.application.knowledge_ingestion_extraction import _parse_chunks
from course_insight.contracts.course import ContentChunk
from course_insight.contracts.errors import DomainError
from course_insight.contracts.intelligence import LLMModelRef
from course_insight.contracts.knowledge_ingestion import (
    KnowledgeCandidate,
    KnowledgeEvidenceRef,
    KnowledgeExtractionResult,
    QuestionConceptLinkCandidate,
)
from course_insight.modules.m0_platform.django_app.models import (
    CourseKnowledgeRelease,
    CourseSource,
    CourseSourceVersion,
    KnowledgeChangeOperation,
    KnowledgeIngestionJob,
    ReleaseConceptSource,
    ReleaseQuestionConceptLink,
    User,
)
from course_insight.modules.m1_course_governance.legacy_powerpoint import (
    LegacyPowerPointConversionUnavailable,
)


pytestmark = pytest.mark.django_db
NOW = datetime(2026, 8, 25, tzinfo=timezone.utc)


class _Extractor:
    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del model_ref, on_retry
        candidates = []
        for index, chunk in enumerate(batch.chunks, start=1):
            if "慢启动" in chunk.text:
                name = "慢启动"
                start = chunk.text.index("慢启动")
            else:
                name = "拥塞控制"
                start = chunk.text.index("拥塞控制")
            candidates.append(
                KnowledgeCandidate(
                    candidate_id=f"{batch.batch_id}-{index}",
                    name=name,
                    description=f"{name}的课程定义",
                    aliases=[],
                    evidence=[
                        KnowledgeEvidenceRef(
                            source_id=chunk.source_id,
                            source_version_id=chunk.source_id,
                            chunk_id=chunk.chunk_id,
                            locator=chunk.locator,
                            span_start=start,
                            span_end=start + len(name),
                            relation_type="definition",
                        )
                    ],
                )
            )
        return KnowledgeExtractionResult(
            result_id=f"result-{batch.batch_id}",
            batch=batch,
            candidates=candidates,
            status="succeeded",
            possibly_truncated=False,
            created_at=created_at,
        )


class _Linker:
    def link(self, question, concepts, *, model_ref, created_at):
        del model_ref, created_at
        return tuple(
            QuestionConceptLinkCandidate(
                question_id=question.question_id,
                concept_id=concept.concept_id,
                confidence=0.9,
                status="usable",
                evidence=list(concept.evidence),
            )
            for concept in concepts
        )


class _OverflowOnceExtractor:
    def __init__(self) -> None:
        self._overflow_returned = False

    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del model_ref, on_retry
        if not self._overflow_returned:
            self._overflow_returned = True
            chunk = batch.chunks[0]
            candidates = [
                KnowledgeCandidate(
                    candidate_id=f"overflow-{index}",
                    name=f"候选知识点 {index}",
                    description="用于触发结果上限后的自动二分重跑。",
                    aliases=[],
                    evidence=[
                        KnowledgeEvidenceRef(
                            source_id=chunk.source_id,
                            source_version_id=chunk.source_id,
                            chunk_id=chunk.chunk_id,
                            locator=chunk.locator,
                            span_start=0,
                            span_end=1,
                            relation_type="mention",
                        )
                    ],
                )
                for index in range(100)
            ]
            return KnowledgeExtractionResult(
                result_id=f"overflow-{batch.batch_id}",
                batch=batch,
                candidates=candidates,
                status="succeeded",
                possibly_truncated=True,
                created_at=created_at,
            )
        return _Extractor().extract(batch, model_ref=None, created_at=created_at)


class _NetworkFailingExtractor:
    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del model_ref, created_at, on_retry
        raise DomainError(
            code="KNOWLEDGE_EXTRACTION_FAILED",
            module="m7",
            message="DeepSeek knowledge extraction did not complete successfully",
            details={
                "batch_id": batch.batch_id,
                "error_code": "DEEPSEEK_NETWORK_ERROR",
            },
            recoverable=True,
        )


class _InvalidParentExtractor:
    def __init__(self) -> None:
        self.batch_sizes: list[int] = []

    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del on_retry
        self.batch_sizes.append(len(batch.chunks))
        if len(batch.chunks) > 1:
            raise DomainError(
                code="KNOWLEDGE_EXTRACTION_OUTPUT_INVALID",
                module="m7",
                message="DeepSeek knowledge extraction output is invalid or ungrounded",
                details={"batch_id": batch.batch_id},
                recoverable=True,
            )
        return _Extractor().extract(
            batch,
            model_ref=model_ref,
            created_at=created_at,
        )


class _IncompleteParentExtractor:
    def __init__(self) -> None:
        self.batch_sizes: list[int] = []

    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del on_retry
        self.batch_sizes.append(len(batch.chunks))
        if len(batch.chunks) > 1:
            raise DomainError(
                code="KNOWLEDGE_EXTRACTION_FAILED",
                module="m7",
                message="DeepSeek knowledge extraction did not complete successfully",
                details={
                    "batch_id": batch.batch_id,
                    "error_code": "DEEPSEEK_INCOMPLETE_RESPONSE",
                },
                recoverable=True,
            )
        return _Extractor().extract(
            batch,
            model_ref=model_ref,
            created_at=created_at,
        )


class _AlwaysIncompleteExtractor:
    def __init__(self) -> None:
        self.call_count = 0

    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del model_ref, created_at, on_retry
        self.call_count += 1
        raise DomainError(
            code="KNOWLEDGE_EXTRACTION_FAILED",
            module="m7",
            message="DeepSeek knowledge extraction did not complete successfully",
            details={
                "batch_id": batch.batch_id,
                "error_code": "DEEPSEEK_INCOMPLETE_RESPONSE",
            },
            recoverable=True,
        )


class _RetryNotifyingExtractor:
    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        if on_retry is None:
            raise AssertionError("knowledge processor did not provide retry observer")
        on_retry(1, 5, "evidence_quote_not_found")
        on_retry(2, 5, "citation_ids_mismatch")
        return _Extractor().extract(
            batch,
            model_ref=model_ref,
            created_at=created_at,
        )


class _ConcurrentExtractor:
    def __init__(self) -> None:
        self._lock = threading.Lock()
        self.active_calls = 0
        self.maximum_active_calls = 0

    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del model_ref, on_retry
        marker = batch.chunks[0].text[0]
        with self._lock:
            self.active_calls += 1
            self.maximum_active_calls = max(
                self.maximum_active_calls,
                self.active_calls,
            )
        try:
            time.sleep({"甲": 0.12, "乙": 0.08, "丙": 0.04}[marker])
            chunk = batch.chunks[0]
            start = chunk.text.index("拥塞控制")
            candidate = KnowledgeCandidate(
                candidate_id=f"candidate-{marker}",
                name=f"{marker}知识点",
                description=f"{marker}批次提取结果",
                aliases=[],
                evidence=[
                    KnowledgeEvidenceRef(
                        source_id=chunk.source_id,
                        source_version_id=chunk.source_id,
                        chunk_id=chunk.chunk_id,
                        locator=chunk.locator,
                        span_start=start,
                        span_end=start + len("拥塞控制"),
                        relation_type="definition",
                    )
                ],
            )
            return KnowledgeExtractionResult(
                result_id=f"result-{marker}",
                batch=batch,
                candidates=[candidate],
                status="succeeded",
                possibly_truncated=False,
                created_at=created_at,
            )
        finally:
            with self._lock:
                self.active_calls -= 1


class _MarkerExtractor:
    def __init__(self, *, failing_markers: frozenset[str] = frozenset()) -> None:
        self.failing_markers = failing_markers
        self.calls: list[str] = []

    def extract(self, batch, *, model_ref, created_at, on_retry=None):
        del model_ref, on_retry
        marker = batch.chunks[0].text[0]
        self.calls.append(marker)
        if marker in self.failing_markers:
            raise DomainError(
                code="KNOWLEDGE_EXTRACTION_FAILED",
                module="m7",
                message="DeepSeek knowledge extraction did not complete successfully",
                details={
                    "batch_id": batch.batch_id,
                    "error_code": "DEEPSEEK_NETWORK_ERROR",
                },
                recoverable=True,
            )
        chunk = batch.chunks[0]
        start = chunk.text.index("拥塞控制")
        candidate = KnowledgeCandidate(
            candidate_id=f"candidate-{marker}",
            name=f"{marker}知识点",
            description=f"{marker}批次提取结果",
            aliases=[],
            evidence=[
                KnowledgeEvidenceRef(
                    source_id=chunk.source_id,
                    source_version_id=chunk.source_id,
                    chunk_id=chunk.chunk_id,
                    locator=chunk.locator,
                    span_start=start,
                    span_end=start + len("拥塞控制"),
                    relation_type="definition",
                )
            ],
        )
        return KnowledgeExtractionResult(
            result_id=f"result-{marker}",
            batch=batch,
            candidates=[candidate],
            status="succeeded",
            possibly_truncated=False,
            created_at=created_at,
        )


def _user() -> User:
    return User.objects.create_user(
        username="pseudonym_pipeline_teacher",
        actor_id="pseudonym_pipeline_teacher",
    )


def _stage_source(
    root: Path,
    user: User,
    name: str,
    text: str | bytes,
    source_type: str = "knowledge",
    *,
    class_id: str = "class_1",
) -> CourseSource:
    source = CourseSource.objects.create(
        course_id="course_1",
        class_id=class_id,
        display_name=name,
        source_type=source_type,
        status="staged",
        created_by=user,
    )
    key = f"{uuid.uuid4().hex}/{uuid.uuid4()}"
    target = root / key
    target.parent.mkdir(parents=True)
    payload = text if isinstance(text, bytes) else text.encode("utf-8")
    target.write_bytes(payload)
    CourseSourceVersion.objects.create(
        source=source,
        version_number=1,
        storage_key=key,
        sha256=hashlib.sha256(payload).hexdigest(),
        media_type=(
            "application/vnd.ms-powerpoint"
            if name.casefold().endswith(".ppt")
            else "text/plain"
        ),
        size_bytes=len(payload),
        status="staged",
    )
    return source


def _job(
    user: User,
    marker: str,
    *,
    class_id: str = "class_1",
) -> KnowledgeIngestionJob:
    return KnowledgeIngestionJob.objects.create(
        course_id="course_1",
        class_id=class_id,
        requested_by=user,
        change_set_checksum=hashlib.sha256(marker.encode()).hexdigest(),
    )


def _delete_job(
    user: User,
    source: CourseSource,
    marker: str,
) -> KnowledgeIngestionJob:
    source.status = "pending_delete"
    source.save(update_fields=("status", "updated_at"))
    job = _job(user, marker)
    KnowledgeChangeOperation.objects.create(
        job=job,
        sequence=1,
        operation=KnowledgeChangeOperation.Operation.DELETE,
        source=source,
        payload={"source_id": str(source.pk)},
    )
    return job


def _upsert_job(
    user: User,
    source: CourseSource,
    marker: str,
) -> KnowledgeIngestionJob:
    job = _job(user, marker)
    version = source.versions.get(status=CourseSourceVersion.Status.STAGED)
    KnowledgeChangeOperation.objects.create(
        job=job,
        sequence=1,
        operation=KnowledgeChangeOperation.Operation.UPSERT,
        source=source,
        payload={
            "source_id": str(source.pk),
            "source_type": source.source_type,
            "version_id": str(version.pk),
            "version_number": version.version_number,
        },
    )
    return job


def _mixed_knowledge_job(
    user: User,
    *,
    deleted: CourseSource,
    added: CourseSource,
    marker: str,
) -> KnowledgeIngestionJob:
    deleted.status = CourseSource.Status.PENDING_DELETE
    deleted.save(update_fields=("status", "updated_at"))
    job = _job(user, marker)
    added_version = added.versions.get(status=CourseSourceVersion.Status.STAGED)
    KnowledgeChangeOperation.objects.bulk_create(
        [
            KnowledgeChangeOperation(
                job=job,
                sequence=1,
                operation=KnowledgeChangeOperation.Operation.DELETE,
                source=deleted,
                payload={"source_id": str(deleted.pk)},
            ),
            KnowledgeChangeOperation(
                job=job,
                sequence=2,
                operation=KnowledgeChangeOperation.Operation.UPSERT,
                source=added,
                payload={
                    "source_id": str(added.pk),
                    "source_type": added.source_type,
                    "version_id": str(added_version.pk),
                    "version_number": added_version.version_number,
                },
            ),
        ]
    )
    return job


def _processor(
    root: Path,
    *,
    legacy_powerpoint_converter=None,
) -> KnowledgeIngestionProcessor:
    return KnowledgeIngestionProcessor(
        storage_root=root,
        extraction_adapter=_Extractor(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        legacy_powerpoint_converter=legacy_powerpoint_converter,
    )


def _pptx_bytes(text: str) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = text
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _fragmented_pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_frame = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(2)
    ).text_frame
    text_frame.paragraphs[0].text = "拥塞控制"
    text_frame.add_paragraph().text = "拥塞控制用于避免网络过载。"
    stream = BytesIO()
    presentation.save(stream)
    return stream.getvalue()


def _three_extraction_chunks() -> list[ContentChunk]:
    chunks = []
    for index, marker in enumerate(("甲", "乙", "丙"), start=1):
        text = marker + "拥塞控制" + ("网" * 5_995)
        chunks.append(
            ContentChunk(
                chunk_id=f"chunk-{marker}",
                source_id=f"source-{marker}",
                text=text,
                locator=f"slide:{index}",
                concept_hints=[],
                sha256=hashlib.sha256(text.encode()).hexdigest(),
            )
        )
    return chunks


def test_two_files_merge_one_concept_without_losing_either_source(tmp_path) -> None:
    user = _user()
    _stage_source(tmp_path, user, "chapter-a.txt", "拥塞控制用于避免网络过载。")
    _stage_source(tmp_path, user, "chapter-b.txt", "拥塞控制会调节发送速率。")
    job = _job(user, "first")

    outcome = _processor(tmp_path).process(job.pk)

    release = CourseKnowledgeRelease.objects.get(pk=outcome.release_id)
    assert release.status == "active"
    assert release.concepts.count() == 1
    assert ReleaseConceptSource.objects.filter(concept__release=release).count() == 2
    assert job.refresh_from_db() is None
    assert job.status == "succeeded"


def test_legacy_powerpoint_is_converted_into_slide_grounded_provenance(
    tmp_path,
) -> None:
    user = _user()
    source = _stage_source(
        tmp_path,
        user,
        "第三章.ppt",
        bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy",
    )
    converted = _pptx_bytes("拥塞控制用于避免网络过载。")

    class _Converter:
        def convert(self, payload: bytes) -> bytes:
            assert payload.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
            return converted

    outcome = _processor(
        tmp_path,
        legacy_powerpoint_converter=_Converter(),
    ).process(_job(user, "legacy-ppt-success").pk)

    evidence = ReleaseConceptSource.objects.get(
        concept__release_id=outcome.release_id,
        source_version__source=source,
    )
    assert outcome.status == KnowledgeIngestionJob.Status.SUCCEEDED
    assert evidence.locator == "slide:1"
    assert evidence.chunk_text == "拥塞控制用于避免网络过载。"


def test_powerpoint_fragments_are_coalesced_into_one_slide_chunk(tmp_path) -> None:
    user = _user()
    payload = _fragmented_pptx_bytes()
    source = _stage_source(tmp_path, user, "fragmented.pptx", payload)
    version = source.versions.get()

    chunks = _parse_chunks(source, version, payload)

    assert len(chunks) == 1
    assert chunks[0].locator == "slide:1"
    assert chunks[0].text == "拥塞控制\n拥塞控制用于避免网络过载。"


def test_question_only_job_reuses_active_knowledge_without_parsing_ppt(
    tmp_path,
) -> None:
    user = _user()
    ppt_source = _stage_source(
        tmp_path,
        user,
        "第三章.ppt",
        bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy",
    )
    _stage_source(
        tmp_path,
        user,
        "existing-questions.txt",
        """[QUESTION]
id: q-existing
type: subjective
stem: 说明拥塞控制的作用。
answer: 避免网络过载。
rubric: 说明避免过载。
explanation: 对应课程原文。
[/QUESTION]""",
        source_type="question",
    )
    converted = _pptx_bytes("拥塞控制用于避免网络过载。")

    class _InitialConverter:
        def convert(self, payload: bytes) -> bytes:
            del payload
            return converted

    first = _processor(
        tmp_path,
        legacy_powerpoint_converter=_InitialConverter(),
    ).process(_job(user, "question-only-baseline").pk)
    baseline = CourseKnowledgeRelease.objects.get(pk=first.release_id)
    assert baseline.concepts.count() == 1

    question_source = _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-question-only
type: fill_blank
stem: 网络过载时使用 ____。
answer: 拥塞控制
explanation: 对应课程原文。
[/QUESTION]""",
        source_type="question",
    )
    job = _job(user, "question-only-change")
    operation = KnowledgeChangeOperation.objects.create(
        job=job,
        sequence=1,
        operation=KnowledgeChangeOperation.Operation.UPSERT,
        source=question_source,
        payload={"source_id": str(question_source.pk)},
    )

    class _NeverExtractKnowledge:
        def extract(self, batch, *, model_ref, created_at, on_retry=None):
            del batch, model_ref, created_at, on_retry
            raise AssertionError("question-only job re-extracted active knowledge")

    class _NeverConvertPowerPoint:
        def convert(self, payload: bytes) -> bytes:
            del payload
            raise AssertionError("question-only job re-parsed the active PowerPoint")

    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_NeverExtractKnowledge(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        legacy_powerpoint_converter=_NeverConvertPowerPoint(),
    )

    outcome = processor.process(job.pk)

    release = CourseKnowledgeRelease.objects.get(pk=outcome.release_id)
    operation.refresh_from_db()
    ppt_source.refresh_from_db()
    assert outcome.status == KnowledgeIngestionJob.Status.SUCCEEDED
    assert release.version_number == baseline.version_number + 1
    assert release.concepts.count() == 1
    assert set(release.questions.values_list("question_id", flat=True)) == {
        "q-existing",
        "q-question-only",
    }
    assert operation.status == KnowledgeChangeOperation.Status.SUCCEEDED
    assert ppt_source.status == CourseSource.Status.ACTIVE


def test_adding_knowledge_reuses_active_provenance_without_reopening_old_file(
    tmp_path,
) -> None:
    user = _user()
    old_source = _stage_source(
        tmp_path,
        user,
        "chapter-old.txt",
        "拥塞控制用于避免网络过载。",
    )
    first = _processor(tmp_path).process(_job(user, "incremental-add-baseline").pk)
    assert CourseKnowledgeRelease.objects.get(pk=first.release_id).concepts.count() == 1

    old_version = old_source.versions.get()
    (tmp_path / old_version.storage_key).unlink()
    new_source = _stage_source(
        tmp_path,
        user,
        "chapter-new.txt",
        "拥塞控制通过调节发送速率缓解网络过载。",
    )

    second = _processor(tmp_path).process(
        _upsert_job(user, new_source, "incremental-add-one-file").pk
    )

    release = CourseKnowledgeRelease.objects.get(pk=second.release_id)
    references = ReleaseConceptSource.objects.filter(concept__release=release)
    assert release.concepts.count() == 1
    assert references.count() == 2
    assert set(references.values_list("source_version__source_id", flat=True)) == {
        old_source.pk,
        new_source.pk,
    }


def test_new_concept_relabels_stored_questions_without_reopening_question_file(
    tmp_path,
) -> None:
    user = _user()
    _stage_source(tmp_path, user, "congestion.txt", "拥塞控制用于避免网络过载。")
    question_source = _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: subjective
stem: 说明拥塞控制和慢启动之间的关系。
answer: 两者都用于网络传输控制。
rubric: 说明两个机制。
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    first = _processor(tmp_path).process(_job(user, "relabel-baseline").pk)
    assert ReleaseQuestionConceptLink.objects.filter(
        question__release_id=first.release_id
    ).count() == 1

    question_version = question_source.versions.get()
    (tmp_path / question_version.storage_key).unlink()
    new_source = _stage_source(
        tmp_path,
        user,
        "slow-start.txt",
        "慢启动逐步增加拥塞窗口。",
    )

    second = _processor(tmp_path).process(
        _upsert_job(user, new_source, "relabel-with-new-concept").pk
    )

    assert set(
        ReleaseQuestionConceptLink.objects.filter(
            question__release_id=second.release_id
        ).values_list("concept__name", flat=True)
    ) == {"拥塞控制", "慢启动"}


def test_mixed_add_and_delete_knowledge_job_remains_incremental(tmp_path) -> None:
    user = _user()
    removed = _stage_source(tmp_path, user, "a.txt", "拥塞控制用于避免网络过载。")
    retained = _stage_source(tmp_path, user, "b.txt", "拥塞控制会调节发送速率。")
    _processor(tmp_path).process(_job(user, "mixed-baseline").pk)
    retained_version = retained.versions.get()
    (tmp_path / retained_version.storage_key).unlink()
    added = _stage_source(
        tmp_path,
        user,
        "c.txt",
        "拥塞控制通过拥塞窗口限制发送速率。",
    )

    outcome = _processor(tmp_path).process(
        _mixed_knowledge_job(
            user,
            deleted=removed,
            added=added,
            marker="mixed-add-delete",
        ).pk
    )

    references = ReleaseConceptSource.objects.filter(
        concept__release_id=outcome.release_id
    )
    assert references.count() == 2
    assert set(references.values_list("source_version__source_id", flat=True)) == {
        retained.pk,
        added.pk,
    }


def test_legacy_powerpoint_failure_is_isolated_with_specific_error_code(
    tmp_path,
) -> None:
    user = _user()
    _stage_source(tmp_path, user, "valid.txt", "拥塞控制用于避免网络过载。")
    legacy = _stage_source(
        tmp_path,
        user,
        "legacy.ppt",
        bytes.fromhex("D0CF11E0A1B11AE1") + b"legacy",
    )
    job = _job(user, "legacy-ppt-unavailable")
    operation = KnowledgeChangeOperation.objects.create(
        job=job,
        sequence=1,
        operation=KnowledgeChangeOperation.Operation.UPSERT,
        source=legacy,
        payload={"source_id": str(legacy.pk)},
    )

    class _UnavailableConverter:
        def convert(self, payload: bytes) -> bytes:
            del payload
            raise LegacyPowerPointConversionUnavailable(
                "Microsoft PowerPoint conversion is unavailable"
            )

    outcome = _processor(
        tmp_path,
        legacy_powerpoint_converter=_UnavailableConverter(),
    ).process(job.pk)

    job.refresh_from_db()
    operation.refresh_from_db()
    assert outcome.status == KnowledgeIngestionJob.Status.PARTIAL
    assert CourseKnowledgeRelease.objects.filter(status="active").count() == 1
    assert job.checkpoint["failed_source_errors"] == {
        str(legacy.pk): "LEGACY_PPT_CONVERSION_UNAVAILABLE"
    }
    assert operation.error_code == "LEGACY_PPT_CONVERSION_UNAVAILABLE"


def test_processing_one_class_never_reads_or_retires_another_class(tmp_path) -> None:
    user = _user()
    _stage_source(
        tmp_path,
        user,
        "class-one.txt",
        "拥塞控制用于避免网络过载。",
        class_id="class_1",
    )
    _stage_source(
        tmp_path,
        user,
        "class-two.txt",
        "慢启动逐步增加拥塞窗口。",
        class_id="class_2",
    )
    processor = _processor(tmp_path)

    first = processor.process(_job(user, "class-one", class_id="class_1").pk)
    second = processor.process(_job(user, "class-two", class_id="class_2").pk)

    first_release = CourseKnowledgeRelease.objects.get(pk=first.release_id)
    second_release = CourseKnowledgeRelease.objects.get(pk=second.release_id)
    assert first_release.status == "active"
    assert second_release.status == "active"
    assert first_release.class_id == "class_1"
    assert second_release.class_id == "class_2"
    assert list(first_release.concepts.values_list("name", flat=True)) == [
        "拥塞控制"
    ]
    assert list(second_release.concepts.values_list("name", flat=True)) == [
        "慢启动"
    ]


def test_bad_file_is_isolated_and_valid_sibling_publishes(tmp_path) -> None:
    user = _user()
    _stage_source(tmp_path, user, "valid.txt", "拥塞控制用于避免网络过载。")
    bad = _stage_source(tmp_path, user, "false.pptx", "不是 PPTX")
    job = _job(user, "partial")
    KnowledgeChangeOperation.objects.create(
        job=job,
        sequence=1,
        operation="upsert_source",
        source=bad,
        payload={"source_id": str(bad.pk)},
    )

    outcome = _processor(tmp_path).process(job.pk)

    assert outcome.status == "partially_succeeded"
    assert outcome.failed_source_ids == (str(bad.pk),)
    assert CourseKnowledgeRelease.objects.filter(status="active").count() == 1
    assert bad.versions.get().status == "failed"
    operation = job.operations.get()
    assert operation.status == "failed"
    assert operation.error_code == "SOURCE_PROCESSING_FAILED"


def test_provider_failure_code_is_preserved_for_teacher_diagnostics(tmp_path) -> None:
    user = _user()
    source = _stage_source(tmp_path, user, "chapter.txt", "拥塞控制用于避免网络过载。")
    job = _job(user, "provider-failure")
    KnowledgeChangeOperation.objects.create(
        job=job,
        sequence=1,
        operation="upsert_source",
        source=source,
        payload={"source_id": str(source.pk)},
    )
    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_NetworkFailingExtractor(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )

    with pytest.raises(DomainError, match="knowledge extraction"):
        processor.process(job.pk)

    job.refresh_from_db()
    operation = job.operations.get()
    assert job.error_code == "DEEPSEEK_NETWORK_ERROR"
    assert job.progress == 25
    assert operation.error_code == "DEEPSEEK_NETWORK_ERROR"


def test_invalid_parent_output_is_retried_as_smaller_batches(tmp_path) -> None:
    user = _user()
    _stage_source(
        tmp_path,
        user,
        "chapter.txt",
        "拥塞控制用于避免网络过载。\n\n慢启动逐步增加拥塞窗口。",
    )
    job = _job(user, "invalid-parent-retry")
    extractor = _InvalidParentExtractor()
    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=extractor,
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )

    outcome = processor.process(job.pk)

    assert outcome.status == KnowledgeIngestionJob.Status.SUCCEEDED
    assert extractor.batch_sizes == [2, 1, 1]


def test_incomplete_parent_response_retries_children_with_partial_progress(
    tmp_path,
) -> None:
    user = _user()
    _stage_source(
        tmp_path,
        user,
        "chapter.txt",
        "拥塞控制用于避免网络过载。\n\n慢启动逐步增加拥塞窗口。",
    )
    job = _job(user, "incomplete-parent-retry")
    extractor = _IncompleteParentExtractor()
    snapshots: list[dict[str, object]] = []

    def heartbeat() -> None:
        job.refresh_from_db()
        snapshots.append(dict(job.checkpoint))

    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=extractor,
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        heartbeat=heartbeat,
    )

    outcome = processor.process(job.pk)

    extraction = [
        checkpoint
        for checkpoint in snapshots
        if checkpoint.get("stage") == "extracting_concepts"
    ]
    total = int(extraction[-1]["total_characters"])
    completed = [int(checkpoint["completed_characters"]) for checkpoint in extraction]
    assert outcome.status == KnowledgeIngestionJob.Status.SUCCEEDED
    assert extractor.batch_sizes == [2, 1, 1]
    assert any(0 < value < total for value in completed)
    assert any(
        checkpoint.get("retry_reason") == "DEEPSEEK_INCOMPLETE_RESPONSE"
        for checkpoint in extraction
    )


def test_incomplete_response_from_minimum_batch_preserves_provider_error(
    tmp_path,
) -> None:
    user = _user()
    _stage_source(tmp_path, user, "short.txt", "拥塞控制用于避免网络过载。")
    job = _job(user, "minimum-incomplete-response")
    extractor = _AlwaysIncompleteExtractor()
    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=extractor,
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )

    with pytest.raises(DomainError, match="did not complete"):
        processor.process(job.pk)

    job.refresh_from_db()
    assert extractor.call_count == 1
    assert job.error_code == "DEEPSEEK_INCOMPLETE_RESPONSE"


def test_output_validation_retries_refresh_activity_without_fake_progress(
    tmp_path,
) -> None:
    user = _user()
    _stage_source(tmp_path, user, "short.txt", "拥塞控制用于避免网络过载。")
    job = _job(user, "output-validation-retry-activity")
    snapshots: list[dict[str, object]] = []

    def heartbeat() -> None:
        job.refresh_from_db()
        snapshots.append(dict(job.checkpoint))

    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_RetryNotifyingExtractor(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        heartbeat=heartbeat,
    )

    outcome = processor.process(job.pk)

    retry_snapshots = [
        checkpoint
        for checkpoint in snapshots
        if checkpoint.get("retry_reason")
        == "KNOWLEDGE_EXTRACTION_OUTPUT_INVALID"
    ]
    assert outcome.status == KnowledgeIngestionJob.Status.SUCCEEDED
    observed_attempts = list(
        dict.fromkeys(checkpoint["retry_attempt"] for checkpoint in retry_snapshots)
    )
    assert observed_attempts == [1, 2]
    assert all(checkpoint["retry_limit"] == 5 for checkpoint in retry_snapshots)
    assert all(checkpoint["completed_characters"] == 0 for checkpoint in retry_snapshots)
    assert retry_snapshots[-1]["validation_code"] == "citation_ids_mismatch"


def test_extraction_uses_three_bounded_workers_and_merges_in_batch_order(
    tmp_path,
) -> None:
    user = _user()
    job = _job(user, "bounded-concurrent-extraction")
    extractor = _ConcurrentExtractor()
    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=extractor,
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )
    candidates, _ = processor._extract_all(
        job,
        "course_1",
        _three_extraction_chunks(),
    )

    assert extractor.maximum_active_calls == 3
    assert [candidate.candidate_id for candidate in candidates] == [
        "candidate-甲",
        "candidate-乙",
        "candidate-丙",
    ]


def test_extraction_resume_skips_successful_batch_checkpoints(tmp_path) -> None:
    user = _user()
    job = _job(user, "resume-extraction-batches")
    first_extractor = _MarkerExtractor(failing_markers=frozenset({"乙"}))
    first_processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=first_extractor,
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        extraction_concurrency=1,
    )

    with pytest.raises(DomainError, match="did not complete"):
        first_processor._extract_all(
            job,
            "course_1",
            _three_extraction_chunks(),
        )

    assert first_extractor.calls == ["甲", "乙"]

    resumed_extractor = _MarkerExtractor()
    resumed_processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=resumed_extractor,
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        extraction_concurrency=1,
    )

    candidates, _ = resumed_processor._extract_all(
        job,
        "course_1",
        _three_extraction_chunks(),
    )

    assert resumed_extractor.calls == ["乙", "丙"]
    assert [candidate.candidate_id for candidate in candidates] == [
        "candidate-甲",
        "candidate-乙",
        "candidate-丙",
    ]
    assert job.extraction_batch_checkpoints.filter(status="succeeded").count() == 3


def test_extraction_failure_records_safe_batch_diagnostics(
    tmp_path,
    caplog,
) -> None:
    class _UnexpectedFailureExtractor:
        def extract(self, batch, *, model_ref, created_at, on_retry=None):
            del batch, model_ref, created_at, on_retry
            raise RuntimeError("private source text must not be logged")

    job = _job(_user(), "safe-extraction-failure-diagnostics")
    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_UnexpectedFailureExtractor(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        extraction_concurrency=1,
    )

    with caplog.at_level(
        logging.ERROR,
        logger="course_insight.application.knowledge_ingestion",
    ), pytest.raises(RuntimeError, match="private source text"):
        processor._extract_all(
            job,
            "course_1",
            _three_extraction_chunks()[:1],
        )

    checkpoint = job.extraction_batch_checkpoints.get()
    assert checkpoint.status == "failed"
    assert checkpoint.batch_index == 1
    assert checkpoint.error_code == "KNOWLEDGE_INGESTION_FAILED"
    assert checkpoint.exception_type == "RuntimeError"
    diagnostic = next(
        record
        for record in caplog.records
        if getattr(record, "_course_insight_event", None)
        == "knowledge_ingestion.extraction_batch_failed"
    )
    assert getattr(diagnostic, "_course_insight_error_code") == (
        "KNOWLEDGE_INGESTION_FAILED"
    )
    assert getattr(diagnostic, "_course_insight_exception_type") == "RuntimeError"
    assert "private source text" not in diagnostic.getMessage()


def test_progress_advances_by_files_characters_and_questions(tmp_path) -> None:
    user = _user()
    paragraph = "拥塞控制用于避免网络过载。"
    _stage_source(
        tmp_path,
        user,
        "chapter-a.txt",
        "\n\n".join(paragraph for _ in range(5)),
    )
    _stage_source(
        tmp_path,
        user,
        "chapter-b.txt",
        "\n\n".join(paragraph for _ in range(5)),
    )
    _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: fill_blank
stem: 网络过载时使用 ____。
answer: 拥塞控制
explanation: 对应课程定义。
[/QUESTION]
[QUESTION]
id: q-2
type: subjective
stem: 说明拥塞控制的作用。
answer: 避免网络过载。
rubric: 说明保护网络路径。
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    job = _job(user, "granular-progress")
    snapshots: list[tuple[int, dict[str, object]]] = []

    def heartbeat() -> None:
        job.refresh_from_db()
        snapshots.append((job.progress, dict(job.checkpoint)))

    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_Extractor(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
        heartbeat=heartbeat,
    )

    processor.process(job.pk)

    reading = [
        checkpoint
        for _, checkpoint in snapshots
        if checkpoint.get("stage") == "reading_sources"
    ]
    extracting = [
        (progress, checkpoint)
        for progress, checkpoint in snapshots
        if checkpoint.get("stage") == "extracting_concepts"
    ]
    linking = [
        checkpoint
        for _, checkpoint in snapshots
        if checkpoint.get("stage") == "linking_questions"
    ]
    assert [checkpoint["completed_units"] for checkpoint in reading] == [0, 1, 2, 3]
    assert len(extracting) >= 2
    assert {checkpoint["total_batches"] for _, checkpoint in extracting} == {1}
    assert extracting[0][1]["completed_characters"] == 0
    assert extracting[-1][1]["completed_characters"] == extracting[-1][1]["total_characters"]
    assert len({progress for progress, _ in extracting}) >= 2
    assert [checkpoint["completed_units"] for checkpoint in linking] == [0, 1, 2]


def test_idempotent_replay_returns_same_release(tmp_path) -> None:
    user = _user()
    _stage_source(tmp_path, user, "chapter.txt", "拥塞控制用于避免网络过载。")
    job = _job(user, "idempotent")
    processor = _processor(tmp_path)

    first = processor.process(job.pk)
    second = processor.process(job.pk)

    assert first.release_id == second.release_id
    assert second.successful_source_ids == first.successful_source_ids
    assert second.failed_source_ids == first.failed_source_ids
    assert second.issue_count == first.issue_count
    assert CourseKnowledgeRelease.objects.count() == 1
    job.refresh_from_db()
    assert job.progress == 100
    assert job.lease_until is None
    assert job.checkpoint["stage"] == "published"


def test_overflow_bisection_publishes_child_chunk_provenance(tmp_path) -> None:
    user = _user()
    text = ("拥塞控制用于避免网络过载。" * 100) + ("慢启动逐步增加拥塞窗口。" * 100)
    _stage_source(tmp_path, user, "long-chapter.txt", text)
    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_OverflowOnceExtractor(),
        question_linking_adapter=_Linker(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )

    outcome = processor.process(_job(user, "overflow-bisection").pk)

    sources = ReleaseConceptSource.objects.filter(concept__release_id=outcome.release_id)
    assert sources.count() >= 2
    assert all(source.chunk_text for source in sources)


def test_delete_rebuild_removes_only_new_release_provenance(tmp_path) -> None:
    user = _user()
    first_source = _stage_source(tmp_path, user, "a.txt", "拥塞控制用于避免网络过载。")
    _stage_source(tmp_path, user, "b.txt", "拥塞控制会调节发送速率。")
    _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: fill_blank
stem: 网络过载时使用 ____。
answer: 拥塞控制
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    first_release_id = _processor(tmp_path).process(_job(user, "before-delete").pk).release_id

    second_release_id = _processor(tmp_path).process(
        _delete_job(user, first_source, "after-delete").pk
    ).release_id

    assert ReleaseConceptSource.objects.filter(concept__release_id=first_release_id).count() == 2
    assert ReleaseConceptSource.objects.filter(concept__release_id=second_release_id).count() == 1
    assert ReleaseQuestionConceptLink.objects.filter(
        question__release_id=second_release_id,
        concept__name="拥塞控制",
    ).count() == 1
    assert CourseKnowledgeRelease.objects.get(pk=first_release_id).status == "retired"
    first_source.refresh_from_db()
    assert first_source.status == "deleted"
    assert set(first_source.versions.values_list("status", flat=True)) == {"deleted"}


def test_deleting_one_knowledge_source_filters_provenance_without_model_calls(
    tmp_path,
) -> None:
    user = _user()
    removed = _stage_source(tmp_path, user, "a.txt", "拥塞控制用于避免网络过载。")
    retained = _stage_source(tmp_path, user, "b.txt", "拥塞控制会调节发送速率。")
    _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: fill_blank
stem: 网络过载时使用 ____。
answer: 拥塞控制
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    _processor(tmp_path).process(_job(user, "local-delete-one-baseline").pk)

    class _NeverExtract:
        def extract(self, batch, *, model_ref, created_at, on_retry=None):
            del batch, model_ref, created_at, on_retry
            raise AssertionError("delete job invoked knowledge extraction")

    class _NeverLink:
        def link(self, question, concepts, *, model_ref, created_at):
            del question, concepts, model_ref, created_at
            raise AssertionError("delete job invoked question linking")

    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_NeverExtract(),
        question_linking_adapter=_NeverLink(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )

    outcome = processor.process(_delete_job(user, removed, "local-delete-one").pk)

    references = ReleaseConceptSource.objects.filter(
        concept__release_id=outcome.release_id
    )
    assert references.count() == 1
    assert references.get().source_version.source_id == retained.pk
    assert ReleaseQuestionConceptLink.objects.filter(
        question__release_id=outcome.release_id,
        concept__name="拥塞控制",
    ).count() == 1


def test_deleting_last_knowledge_source_removes_concept_without_model_calls(
    tmp_path,
) -> None:
    user = _user()
    removed = _stage_source(
        tmp_path,
        user,
        "only.txt",
        "拥塞控制用于避免网络过载。",
    )
    _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: fill_blank
stem: 网络过载时使用 ____。
answer: 拥塞控制
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    _processor(tmp_path).process(_job(user, "local-delete-last-baseline").pk)

    class _NeverExtract:
        def extract(self, batch, *, model_ref, created_at, on_retry=None):
            del batch, model_ref, created_at, on_retry
            raise AssertionError("delete job invoked knowledge extraction")

    class _NeverLink:
        def link(self, question, concepts, *, model_ref, created_at):
            del question, concepts, model_ref, created_at
            raise AssertionError("delete job invoked question linking")

    processor = KnowledgeIngestionProcessor(
        storage_root=tmp_path,
        extraction_adapter=_NeverExtract(),
        question_linking_adapter=_NeverLink(),
        model_ref=LLMModelRef(
            model_name="deepseek-v4-flash",
            model_version="runtime-api",
            status="configured",
        ),
        clock=lambda: NOW,
    )

    outcome = processor.process(_delete_job(user, removed, "local-delete-last").pk)

    release = CourseKnowledgeRelease.objects.get(pk=outcome.release_id)
    assert release.concepts.count() == 0
    assert release.questions.count() == 1
    assert ReleaseQuestionConceptLink.objects.filter(
        question__release=release
    ).count() == 0


def test_question_concept_set_shrinks_when_a_concept_loses_its_last_source(
    tmp_path,
) -> None:
    user = _user()
    _stage_source(tmp_path, user, "congestion.txt", "拥塞控制用于避免网络过载。")
    slow_start = _stage_source(tmp_path, user, "slow-start.txt", "慢启动逐步增加拥塞窗口。")
    _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: subjective
stem: 说明拥塞控制与慢启动。
answer: 两者都是网络传输机制。
rubric: 说明两个知识点。
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    processor = _processor(tmp_path)
    first = processor.process(_job(user, "two-question-concepts").pk)

    second = processor.process(
        _delete_job(user, slow_start, "remove-one-question-concept").pk
    )

    assert set(
        ReleaseQuestionConceptLink.objects.filter(
            question__release_id=first.release_id
        ).values_list("concept__name", flat=True)
    ) == {"拥塞控制", "慢启动"}
    assert set(
        ReleaseQuestionConceptLink.objects.filter(
            question__release_id=second.release_id
        ).values_list("concept__name", flat=True)
    ) == {"拥塞控制"}


def test_all_questions_are_relinked_against_new_concept_set(tmp_path) -> None:
    user = _user()
    _stage_source(tmp_path, user, "chapter.txt", "拥塞控制用于避免网络过载。")
    _stage_source(
        tmp_path,
        user,
        "questions.txt",
        """[QUESTION]
id: q-1
type: fill_blank
stem: 网络过载时使用 ____。
answer: 拥塞控制
explanation: 对应课程定义。
[/QUESTION]""",
        source_type="question",
    )
    first = _processor(tmp_path).process(_job(user, "questions-one").pk)
    _stage_source(tmp_path, user, "slow-start.txt", "慢启动逐步增加拥塞窗口。")

    second = _processor(tmp_path).process(_job(user, "questions-two").pk)

    assert ReleaseQuestionConceptLink.objects.filter(question__release_id=first.release_id).count() == 1
    assert ReleaseQuestionConceptLink.objects.filter(question__release_id=second.release_id).count() == 2
